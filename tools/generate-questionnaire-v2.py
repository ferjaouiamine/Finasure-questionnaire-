import json
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "Questionnaire_Maturite_ERM_V2.xlsx"
TARGET = ROOT / "js" / "questionnaire-data.js"


def slug(value):
    value = unicodedata.normalize("NFD", value)
    value = "".join(char for char in value if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


workbook = openpyxl.load_workbook(SOURCE, data_only=False)
questionnaire = workbook["1. Questionnaire"]
notice = workbook["0. Notice"]
recommendations_sheet = workbook["4. Référentiel recos"]

weights = {str(notice.cell(row, 2).value): int(notice.cell(row, 3).value) for row in range(41, 52)}
dimension_rows = []
current_dimension = None
questions = []

answer_definitions = [
    {"code": "A", "name": "Non", "score": 1, "points": 0, "description": "La pratique décrite n’existe pas, ou seulement de façon informelle et non reproductible."},
    {"code": "B", "name": "Partiellement", "score": 2, "points": 1, "description": "La pratique existe sur une partie du périmètre, de façon irrégulière ou sans preuve d’efficacité."},
    {"code": "C", "name": "Oui", "score": 3, "points": 2, "description": "La pratique est appliquée régulièrement sur tout le périmètre et son efficacité est démontrable."},
]

for row in range(5, 27):
    if questionnaire.cell(row, 1).value:
        current_dimension = str(questionnaire.cell(row, 1).value).strip()
        dimension_rows.append(current_dimension)
    number = str(questionnaire.cell(row, 2).value).strip()
    text = str(questionnaire.cell(row, 3).value).strip()
    dimension_name = re.sub(r"^\d+\.\s*", "", current_dimension)
    questions.append({
        "id": len(questions) + 1,
        "reference": number,
        "dimensionId": slug(dimension_name),
        "text": text,
        "answers": answer_definitions,
    })

recommendations = {}
for row in range(4, recommendations_sheet.max_row + 1):
    dimension = recommendations_sheet.cell(row, 2).value
    level = recommendations_sheet.cell(row, 3).value
    if not dimension or not level:
        continue
    dimension_name = re.sub(r"^\d+\.\s*", "", str(dimension).strip())
    recommendations.setdefault(slug(dimension_name), {})[str(level).strip()] = {
        "diagnostic": str(recommendations_sheet.cell(row, 4).value or "").strip(),
        "shortTerm": str(recommendations_sheet.cell(row, 5).value or "").strip(),
        "mediumTerm": str(recommendations_sheet.cell(row, 6).value or "").strip(),
    }

dimensions = []
for order, full_name in enumerate(dimension_rows):
    name = re.sub(r"^\d+\.\s*", "", full_name)
    dimensions.append({
        "id": slug(name),
        "name": name,
        "weight": weights[full_name],
        "order": order,
        "recommendations": recommendations.get(slug(name), {}),
    })

payload = {
    "version": "2026.2-v2",
    "sourceStatus": "Questionnaire_Maturite_ERM_V2.xlsx",
    "answerScale": {"storedMin": 1, "storedMax": 3, "pointsMin": 0, "pointsMax": 2},
    "dimensions": dimensions,
    "questions": questions,
    "steps": [[1, 6], [7, 12], [13, 18], [19, 22]],
    "stepNames": ["Fondations", "Maîtrise", "Pilotage", "Résilience"],
    "contactUrl": "https://www.finasure-solutions.com/contact/",
}

content = "(function(){\n  \"use strict\";\n  window.FINASURE_ERM_DATA = Object.freeze(" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ");\n})();\n"
TARGET.write_text(content, encoding="utf-8")
print(f"Generated {TARGET} with {len(questions)} questions and {len(dimensions)} dimensions")
